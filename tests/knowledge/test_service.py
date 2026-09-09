import asyncio
from pathlib import Path
import pytest
from src.knowledge.service import KnowledgeBaseService
from src.knowledge.parsers import chunk_blocks, extract


@pytest.fixture
def service(tmp_path):
    return KnowledgeBaseService(tmp_path)


async def add(service, library, text, **metadata):
    result = service.enqueue(library, 'report.md', text.encode(), metadata)
    if result['job_id']:
        await service.process(result['job_id'])
        assert service.job(result['job_id'])['status'] == 'ready'
    return result


@pytest.mark.asyncio
async def test_ingest_search_reopen_and_duplicate(service):
    kb = service.create_library('公司研究')['id']
    first = await add(service, kb, '# 毛利率\n\n毛利率下降，主要原因是原材料成本上涨。', ticker='600001', market='A', published_at='2025-03-01')
    second = await add(service, kb, '# 毛利率\n\n毛利率下降，主要原因是原材料成本上涨。', ticker='600001', market='A', published_at='2025-03-01')
    assert second['duplicate'] and first['doc_id'] == second['doc_id']
    reopened = KnowledgeBaseService(service.root)
    result = await reopened.search('毛利率 成本', [kb])
    assert '原材料' in result['results'][0]['text']
    assert 'paragraph' in result['results'][0]['locator']
    assert result['mode'] == 'keyword'


@pytest.mark.asyncio
async def test_filter_before_ranking_no_future_unknown_or_other_library(service):
    kb = service.create_library('A')['id']; other = service.create_library('B')['id']
    await add(service, kb, 'margin known published facts', ticker='A', published_at='2025-01-01')
    await add(service, kb, 'margin margin future facts', ticker='A', published_at='2026-01-01')
    await add(service, kb, 'margin unknown date', ticker='A')
    await add(service, other, 'margin other library', ticker='A', published_at='2025-01-01')
    await add(service, kb, 'margin wrong company', ticker='B', published_at='2025-01-01')
    result = await service.search('margin', [kb], {'ticker': 'A', 'as_of': '2025-12-31'})
    assert [r['text'] for r in result['results']] == ['margin known published facts']


@pytest.mark.asyncio
async def test_deleted_live_but_historical_snapshot_reproducible(service):
    kb = service.create_library('A')['id']
    doc = await add(service, kb, 'retained evidence margin')
    snapshot = service.corpus([kb])
    service.delete_document(doc['doc_id'])
    assert not (await service.search('margin', [kb]))['results']
    assert (await service.search('margin', snapshot=snapshot))['results']
    with pytest.raises(KeyError):
        service.original_path(doc['doc_id'])
    new = await add(service, kb, 'retained evidence margin')
    assert new['doc_id'] != doc['doc_id']


@pytest.mark.asyncio
async def test_failed_and_restart_recovery(service):
    kb = service.create_library('A')['id']
    entry = service.enqueue(kb, 'broken.pdf', b'not pdf')
    await service.process(entry['job_id'])
    assert service.job(entry['job_id'])['status'] == 'failed'
    assert not service.corpus([kb])
    entry = service.enqueue(kb, 'ok.txt', b'hello evidence')
    service._status(entry['job_id'], 'parsing')
    recovered = service.recover_jobs()
    assert entry['job_id'] in recovered
    await service.process(entry['job_id'])
    assert service.job(entry['job_id'])['status'] == 'ready'


@pytest.mark.asyncio
async def test_metadata_changes_create_immutable_versions_and_ai_exclusion(service):
    kb = service.create_library('A')['id']
    a = await add(service, kb, 'source evidence', published_at='2025-01-01')
    b = await add(service, kb, 'source evidence', published_at='2025-02-01')
    await add(service, kb, 'AI evidence', source_type='ai_generated')
    assert a['doc_id'] != b['doc_id']
    assert len(service.corpus([kb])) == 2


class Embedding:
    def __init__(self): self.calls = []
    async def generate_embeddings(self, texts):
        self.calls.extend(texts)
        return [[1.0, float('margin' in t)] for t in texts]


@pytest.mark.asyncio
async def test_embedding_cache_hybrid_and_model_version(service):
    model = Embedding(); service.embedding = model; service.model_id = 'test:v1'
    kb = service.create_library('A')['id']
    await add(service, kb, 'margin rose')
    first = await service.search('profit margin', [kb])
    count = len(model.calls)
    await service.search('profit margin', [kb])
    assert len(model.calls) == count
    assert first['mode'] == 'hybrid'
    service.model_id = 'test:v2'
    await service.search('profit margin', [kb])
    assert len(model.calls) > count


@pytest.mark.asyncio
async def test_embedding_failure_falls_back(service):
    kb = service.create_library('A')['id']
    await add(service, kb, 'margin fell')
    class Broken:
        async def generate_embeddings(self, texts): raise RuntimeError('offline')
    service.embedding = Broken(); service.model_id = 'broken'
    result = await service.search('margin', [kb])
    assert result['results'] and result['mode'] == 'keyword'
    assert 'fallback' in result['warnings'][0]


@pytest.mark.asyncio
async def test_delete_during_embedding_never_returns_document(service):
    kb = service.create_library('A')['id']; doc = await add(service, kb, 'margin facts')
    class Deleting:
        async def generate_embeddings(self, texts):
            if service.documents(kb): service.delete_document(doc['doc_id'])
            return [[1, 0] for _ in texts]
    service.embedding = Deleting(); service.model_id = 'test'
    assert not (await service.search('margin', [kb]))['results']


def test_upload_validation_and_path_independence(service):
    kb = service.create_library('A')['id']
    with pytest.raises(ValueError): service.enqueue(kb, 'file.exe', b'x')
    with pytest.raises(ValueError): service.enqueue(kb, 'file.txt', b'x', {'published_at': 'bad'})
    result = service.enqueue(kb, '../../external.txt', b'valid')
    assert service.original_path(result['doc_id']).parent == service.root / 'originals'
    assert service.document(result['doc_id'])['filename'] == 'external.txt'


def test_structure_and_table_header_retention(tmp_path):
    from docx import Document
    doc = Document(); doc.add_heading('Financials', 1); doc.add_paragraph('Margin rose.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Period'; table.cell(0, 1).text = 'Revenue'
    table.cell(1, 0).text = '2025'; table.cell(1, 1).text = '100'
    path = tmp_path / 'report.docx'; doc.save(path)
    blocks = extract(path)
    assert any(b['kind'] == 'table' and '100' in b['text'] for b in blocks)
    assert all(b['page'] is None for b in blocks)
    chunks = list(chunk_blocks([{'text': 'Year | Revenue\n' + '\n'.join('2025 | 100' for _ in range(20)), 'kind': 'table', 'locator': 'table 1', 'page': 1}], max_chars=70))
    assert len(chunks) > 1 and all(c['text'].startswith('Year | Revenue') for c in chunks)


@pytest.mark.asyncio
async def test_qdrant_adapter_with_embedded_test_server(service, monkeypatch):
    from qdrant_client import QdrantClient, models
    client = QdrantClient(':memory:')
    client.create_collection('test', vectors_config=models.VectorParams(size=2, distance=models.Distance.COSINE))
    monkeypatch.setattr(client, 'close', lambda: None)
    monkeypatch.setattr(service, '_qdrant', lambda dimension: (client, 'test'))
    service.embedding = Embedding(); service.model_id = 'test'; service.backend = 'qdrant'
    kb = service.create_library('A')['id']; other = service.create_library('B')['id']
    await add(service, kb, 'margin revenue')
    await add(service, other, 'margin other')
    result = await service.search('margin', [kb])
    assert result['mode'] == 'hybrid' and all(r['kb_id'] == kb for r in result['results'])


def make_pdf(text):
    # Minimal valid one-page fixture with a standard font; avoids renderer dependencies.
    stream = f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'.encode()
    bodies = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
              b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
              b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>', b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    output = b'%PDF-1.4\n'; offsets = [0]
    for i, body in enumerate(bodies, 1):
        offsets.append(len(output)); output += f'{i} 0 obj\n'.encode() + body + b'\nendobj\n'
    xref = len(output)
    output += b'xref\n0 6\n0000000000 65535 f \n'
    output += b''.join(f'{offset:010} 00000 n \n'.encode() for offset in offsets[1:])
    return output + f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()


@pytest.mark.asyncio
async def test_pdf_page_locator_and_empty_scan_status(service):
    kb = service.create_library('A')['id']
    entry = service.enqueue(kb, 'annual.pdf', make_pdf('Revenue increased'))
    await service.process(entry['job_id'])
    result = await service.search('Revenue', [kb])
    assert result['results'][0]['page'] == 1
    assert result['results'][0]['locator'] == 'PDF p.1'
    blank = service.enqueue(kb, 'scan.pdf', make_pdf(''))
    await service.process(blank['job_id'])
    assert service.job(blank['job_id'])['status'] == 'needs_ocr'
